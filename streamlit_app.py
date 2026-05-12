"""
English Vocabulary Extractor — Streamlit App.
Deploy on Streamlit Cloud for free live hosting.
"""

import os
import io
import time
import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# ── Set API key from Streamlit Cloud secrets before importing src modules ──
try:
    if "OPENAI_API_KEY" in st.secrets:
        os.environ["OPENAI_API_KEY"] = st.secrets["OPENAI_API_KEY"]
except FileNotFoundError:
    pass  # Running locally — .env will be used instead

from src.transcript import get_transcript
from src.vocabulary import extract_vocabulary
from src.chat import chat_about_video

# ── Page Config ───────────────────────────────────────────────────
st.set_page_config(
    page_title="English Vocabulary Extractor | AI-Powered",
    page_icon="🎓",
    layout="wide",
    initial_sidebar_state="collapsed",
)

# ── Custom CSS ────────────────────────────────────────────────────
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700;800&display=swap');

.stApp { font-family: 'Inter', sans-serif; }
#MainMenu, footer, header { visibility: hidden; }

.hero-title {
    font-size: 2.6rem; font-weight: 800; text-align: center; margin-bottom: 0;
    background: linear-gradient(135deg, #818CF8, #6366F1, #A78BFA);
    -webkit-background-clip: text; -webkit-text-fill-color: transparent;
}
.hero-sub { text-align: center; color: #9CA3AF; font-size: 1.05rem; margin-bottom: 1.5rem; }
.gradient-divider { height: 2px; background: linear-gradient(90deg, transparent, #6366F1, transparent); margin: 1.5rem 0; border: none; }

.stat-card {
    background: linear-gradient(135deg, rgba(99,102,241,0.12), rgba(139,92,246,0.06));
    border: 1px solid rgba(99,102,241,0.25); border-radius: 12px;
    padding: 1rem; text-align: center;
}
.stat-value { font-size: 1.6rem; font-weight: 700; color: #818CF8; }
.stat-label { font-size: 0.82rem; color: #9CA3AF; margin-top: 0.2rem; }

.badge-beginner    { background: rgba(34,197,94,0.18); color: #22C55E; padding: 2px 10px; border-radius: 20px; font-size: 0.78rem; font-weight: 600; }
.badge-intermediate{ background: rgba(234,179,8,0.18); color: #EAB308; padding: 2px 10px; border-radius: 20px; font-size: 0.78rem; font-weight: 600; }
.badge-advanced    { background: rgba(239,68,68,0.18); color: #EF4444; padding: 2px 10px; border-radius: 20px; font-size: 0.78rem; font-weight: 600; }

.urdu-text { direction: rtl; text-align: right; font-size: 1.05rem; color: #A78BFA; }

.stDownloadButton > button {
    background: rgba(99,102,241,0.12) !important; border: 1px solid rgba(99,102,241,0.3) !important;
    border-radius: 10px !important; font-weight: 500 !important;
}
.stDownloadButton > button:hover { background: rgba(99,102,241,0.22) !important; }
</style>
""", unsafe_allow_html=True)

# ── Session State ─────────────────────────────────────────────────
for key, val in {
    "vocabulary": None, "transcript": None, "movie_title": None,
    "video_id": None, "lang": None, "chat_messages": [],
}.items():
    if key not in st.session_state:
        st.session_state[key] = val


# ── Helpers ───────────────────────────────────────────────────────
def make_excel(vocab_list, movie_title, full=True):
    """Create an Excel file in memory and return bytes."""
    buf = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active

    if full:
        ws.title = "Vocabulary"
        headers = ["English Word", "Urdu Meaning", "English Meaning", "Example Sentence", "Level"]
        widths = [22, 28, 32, 50, 15]
        fields = ["word", "urdu_meaning", "english_meaning", "example_sentence", "level"]
    else:
        ws.title = "English_Urdu"
        headers = ["English Word", "Urdu Meaning"]
        widths = [22, 32]
        fields = ["word", "urdu_meaning"]

    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="6366F1", end_color="6366F1", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    hc = ws.cell(row=2, column=1, value=f"🎬 Movie/Video: {movie_title}")
    hc.font = Font(bold=True, size=14, color="FFFFFF")
    hc.fill = PatternFill(start_color="33334C", end_color="33334C", fill_type="solid")
    hc.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 30

    for i, item in enumerate(vocab_list, 3):
        for c, f in enumerate(fields, 1):
            ws.cell(row=i, column=c, value=item.get(f, ""))

    for c, w in enumerate(widths):
        ws.column_dimensions[chr(65 + c)].width = w

    wb.save(buf)
    return buf.getvalue()


def badge(level):
    lvl = (level or "intermediate").lower()
    cls = f"badge-{lvl}" if lvl in ("beginner", "intermediate", "advanced") else "badge-intermediate"
    return f'<span class="{cls}">{level}</span>'


# ══════════════════════════════════════════════════════════════════
#  UI
# ══════════════════════════════════════════════════════════════════

st.markdown('<h1 class="hero-title">🎓 English Vocabulary Extractor</h1>', unsafe_allow_html=True)
st.markdown('<p class="hero-sub">YouTube → AI → Your English Vocabulary 🚀&nbsp;&nbsp;|&nbsp;&nbsp;Powered by GPT-4o · For Urdu-speaking learners</p>', unsafe_allow_html=True)
st.markdown('<div class="gradient-divider"></div>', unsafe_allow_html=True)

# ── URL Input ─────────────────────────────────────────────────────
c1, c2 = st.columns([5, 1])
with c1:
    url = st.text_input("YouTube URL", placeholder="https://www.youtube.com/watch?v=...", label_visibility="collapsed")
with c2:
    go = st.button("🚀 Extract", type="primary", use_container_width=True)

# ── Extraction ────────────────────────────────────────────────────
if go and url:
    try:
        prog = st.progress(0, text="🔍 Fetching transcript from YouTube...")
        tdata = get_transcript(url)
        prog.progress(40, text="✅ Transcript ready — AI is analyzing vocabulary...")

        result = extract_vocabulary(transcript_text=tdata["full_text"], video_title=None)
        prog.progress(100, text="🎉 Done!")

        st.session_state.vocabulary = result.get("vocabulary", [])
        st.session_state.movie_title = result.get("movie_title", "Unknown")
        st.session_state.transcript = tdata["full_text"]
        st.session_state.video_id = tdata["video_id"]
        st.session_state.lang = tdata["language"]
        st.session_state.chat_messages = []

        time.sleep(0.8)
        prog.empty()
        st.rerun()
    except ValueError as e:
        st.error(f"⚠️ {e}")
    except Exception as e:
        st.error(f"❌ Error: {e}")
elif go and not url:
    st.warning("Please paste a YouTube URL first.")

# ── Results ───────────────────────────────────────────────────────
if st.session_state.vocabulary:
    vocab = st.session_state.vocabulary
    title = st.session_state.movie_title

    st.markdown('<div class="gradient-divider"></div>', unsafe_allow_html=True)

    # Stats
    s1, s2, s3, s4 = st.columns(4)
    with s1:
        st.markdown(f'<div class="stat-card"><div class="stat-value">🎬</div><div class="stat-label">{title}</div></div>', unsafe_allow_html=True)
    with s2:
        st.markdown(f'<div class="stat-card"><div class="stat-value">{len(vocab)}</div><div class="stat-label">Words Extracted</div></div>', unsafe_allow_html=True)
    with s3:
        counts = {"beginner": 0, "intermediate": 0, "advanced": 0}
        for v in vocab:
            lvl = (v.get("level", "")).lower()
            if lvl in counts:
                counts[lvl] += 1
        st.markdown(f'<div class="stat-card"><div class="stat-value">{counts["beginner"]}/{counts["intermediate"]}/{counts["advanced"]}</div><div class="stat-label">Beginner / Inter / Adv</div></div>', unsafe_allow_html=True)
    with s4:
        st.markdown(f'<div class="stat-card"><div class="stat-value">{(st.session_state.lang or "en").upper()}</div><div class="stat-label">Transcript Language</div></div>', unsafe_allow_html=True)

    st.markdown("")

    # ── Filters ───────────────────────────────────────────────────
    f1, f2 = st.columns([1, 3])
    with f1:
        level_filter = st.selectbox("Filter by Level", ["All", "Beginner", "Intermediate", "Advanced"])
    with f2:
        search = st.text_input("🔍 Search", placeholder="Search words...", label_visibility="collapsed")

    filtered = vocab
    if level_filter != "All":
        filtered = [v for v in filtered if (v.get("level", "")).lower() == level_filter.lower()]
    if search:
        q = search.lower()
        filtered = [v for v in filtered if q in (v.get("word", "")).lower() or q in (v.get("english_meaning", "")).lower() or q in (v.get("urdu_meaning", "")).lower()]

    st.caption(f"Showing **{len(filtered)}** of {len(vocab)} words")

    # ── Vocabulary Table ──────────────────────────────────────────
    if filtered:
        df = pd.DataFrame([{
            "Word": v.get("word", ""),
            "Urdu Meaning": v.get("urdu_meaning", ""),
            "English Meaning": v.get("english_meaning", ""),
            "Example": v.get("example_sentence", ""),
            "Level": v.get("level", ""),
        } for v in filtered])

        st.dataframe(
            df,
            use_container_width=True,
            hide_index=True,
            height=min(len(df) * 38 + 40, 600),
            column_config={
                "Word": st.column_config.TextColumn("Word", width="medium"),
                "Urdu Meaning": st.column_config.TextColumn("اردو معنی", width="medium"),
                "English Meaning": st.column_config.TextColumn("Meaning", width="large"),
                "Example": st.column_config.TextColumn("Example", width="large"),
                "Level": st.column_config.TextColumn("Level", width="small"),
            },
        )
    else:
        st.info("No words match your filter.")

    # ── Downloads ─────────────────────────────────────────────────
    st.markdown('<div class="gradient-divider"></div>', unsafe_allow_html=True)
    d1, d2, d3 = st.columns(3)
    with d1:
        st.download_button(
            "📥 Full Excel", make_excel(vocab, title, full=True),
            file_name=f"{title}_vocabulary.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )
    with d2:
        st.download_button(
            "📥 English + Urdu Only", make_excel(vocab, title, full=False),
            file_name=f"{title}_english_urdu.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )
    with d3:
        csv = pd.DataFrame(vocab).to_csv(index=False)
        st.download_button(
            "📥 CSV", csv, file_name=f"{title}_vocabulary.csv",
            mime="text/csv", use_container_width=True,
        )

    # ── Chat ──────────────────────────────────────────────────────
    st.markdown('<div class="gradient-divider"></div>', unsafe_allow_html=True)
    st.markdown("### 💬 Ask About This Video")
    st.caption("Ask any question — the AI tutor answers using only the video transcript.")

    for msg in st.session_state.chat_messages:
        with st.chat_message(msg["role"]):
            st.markdown(msg["content"])

    if prompt := st.chat_input("Ask a question about the video..."):
        st.session_state.chat_messages.append({"role": "user", "content": prompt})
        with st.chat_message("user"):
            st.markdown(prompt)

        with st.chat_message("assistant"):
            with st.spinner("Thinking..."):
                reply = chat_about_video(
                    title=st.session_state.movie_title,
                    transcript=st.session_state.transcript,
                    conversation=st.session_state.chat_messages,
                )
            st.markdown(reply)
        st.session_state.chat_messages.append({"role": "assistant", "content": reply})

# ── Footer ────────────────────────────────────────────────────────
st.markdown('<div class="gradient-divider"></div>', unsafe_allow_html=True)
st.markdown(
    '<p style="text-align:center; color:#4B5563; font-size:0.85rem;">'
    'Built with ❤️ for Urdu-speaking English learners&nbsp;&nbsp;|&nbsp;&nbsp;'
    'Powered by OpenAI GPT-4o &amp; Streamlit</p>',
    unsafe_allow_html=True,
)
