import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl
import pytest

from src import excel

SAMPLE_DATA = {
    "movie_title": "Test Movie",
    "vocabulary": [
        {"word": "devastated", "urdu_meaning": "تباہ حال", "english_meaning": "extremely shocked",
         "example_sentence": "I was devastated.", "level": "Advanced"},
        {"word": "shrink", "urdu_meaning": "سکڑنا", "english_meaning": "become smaller",
         "example_sentence": "It shrinks in the wash.", "level": "Intermediate"},
        {"word": "go beyond", "urdu_meaning": "آگے بڑھنا", "english_meaning": "exceed",
         "example_sentence": "Go beyond your limits.", "level": "Advanced"},
    ],
}


@pytest.fixture(autouse=True)
def isolated_file(tmp_path, monkeypatch):
    monkeypatch.setattr(excel, "EXCEL_FILE", str(tmp_path / "English_tutor.xlsx"))


def test_first_save_writes_all():
    r = excel.save_to_excel(SAMPLE_DATA)
    assert r["saved"] == 3
    assert r["skipped"] == 0
    assert os.path.exists(excel.EXCEL_FILE)


def test_duplicate_save_skips_all():
    excel.save_to_excel(SAMPLE_DATA)
    r = excel.save_to_excel(SAMPLE_DATA)
    assert r["saved"] == 0
    assert r["skipped"] == 3
    assert r["duplicates"] == ["devastated", "shrink", "go beyond"]


def test_partial_duplicates():
    excel.save_to_excel(SAMPLE_DATA)
    r = excel.save_to_excel({
        "movie_title": "Movie 2",
        "vocabulary": [
            {"word": "shrink", "urdu_meaning": "x", "english_meaning": "y",
             "example_sentence": "z", "level": "Beginner"},
            {"word": "shed", "urdu_meaning": "بہانا", "english_meaning": "let fall",
             "example_sentence": "Trees shed leaves.", "level": "Advanced"},
        ],
    })
    assert r["saved"] == 1
    assert r["skipped"] == 1
    assert r["duplicates"] == ["shrink"]


def test_movie_heading_written():
    excel.save_to_excel(SAMPLE_DATA)
    wb = openpyxl.load_workbook(excel.EXCEL_FILE)
    ws = wb.active
    assert ws.cell(row=2, column=1).value == "🎬 Movie/Video: Test Movie"
    # 1 header + 1 heading + 3 data rows
    assert ws.max_row == 5


def test_english_urdu_sheet_dedup_per_sheet():
    excel.save_to_excel(SAMPLE_DATA)
    r = excel.save_english_urdu_excel({
        "movie_title": "Test Movie",
        "vocabulary": [
            {"word": "devastated", "urdu_meaning": "تباہ حال"},
            {"word": "brilliant", "urdu_meaning": "ذہین"},
        ],
    })
    # Both are new to the English_Urdu sheet, so both are saved
    assert r["saved"] == 2
    wb = openpyxl.load_workbook(excel.EXCEL_FILE)
    assert "English_Urdu" in wb.sheetnames

    # Second en-ur save with the same words skips them
    r2 = excel.save_english_urdu_excel({
        "movie_title": "Test Movie 2",
        "vocabulary": [{"word": "brilliant", "urdu_meaning": "ذہین"}],
    })
    assert r2["saved"] == 0
    assert r2["skipped"] == 1